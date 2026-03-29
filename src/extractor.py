import os


class UnsupportedFormatError(Exception):
    pass


def extract_text(path: str) -> str:
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".md":
        return _extract_md(path)
    elif ext == ".pdf":
        return _extract_pdf(path)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext == ".xlsx":
        return _extract_xlsx(path)
    elif ext == ".bib":
        return _extract_bib(path)
    else:
        raise UnsupportedFormatError(f"Unsupported format: {ext}")


def _extract_md(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_pdf(path: str) -> str:
    import fitz  # pymupdf
    doc = fitz.open(path)
    return "\n".join(page.get_text() for page in doc)


def _extract_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)


def _extract_bib(path: str) -> str:
    import bibtexparser
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        library = bibtexparser.load(f)
    entries = []
    for entry in library.entries:
        parts = [f"[{entry.get('ENTRYTYPE', 'ref').upper()}] {entry.get('ID', '')}"]
        for field in ["title", "author", "year", "abstract", "journal", "booktitle"]:
            if field in entry:
                parts.append(f"{field}: {entry[field]}")
        entries.append("\n".join(parts))
    return "\n\n".join(entries)
